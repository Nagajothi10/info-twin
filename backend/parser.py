import os
import pymupdf
import docx
import pandas as pd

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt', 'csv', 'xlsx'}

def is_allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_content(file_path, filename):
    """
    Extracts real text and structured content from supported files:
    1. PDF -> PyMuPDF (pymupdf)
    2. DOCX -> python-docx
    3. TXT -> normal text reading
    4. CSV -> pandas
    5. XLSX -> pandas
    """
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else 'txt'
    file_size_bytes = os.path.getsize(file_path) if os.path.exists(file_path) else 0

    if file_size_bytes < 1024:
        size_str = f"{file_size_bytes} B"
    elif file_size_bytes < 1024 * 1024:
        size_str = f"{(file_size_bytes / 1024):.1f} KB"
    else:
        size_str = f"{(file_size_bytes / (1024 * 1024)):.1f} MB"

    content = ""
    summary = f"Processed {ext.upper()} document ({size_str})."
    extracted_claim = f"Extracted text content from {filename}."

    try:
        # 1. PDF -> PyMuPDF
        if ext == 'pdf':
            doc = pymupdf.open(file_path)
            total_pages = len(doc)
            pages_text = []
            for page_index in range(total_pages):
                page = doc[page_index]
                text = page.get_text()
                if text.strip():
                    pages_text.append(f"--- Page {page_index + 1} ---\n{text.strip()}")
            content = "\n\n".join(pages_text) if pages_text else "(No extractable text found in PDF document)"
            summary = f"PyMuPDF extracted {len(content)} characters across {total_pages} page(s)."
            first_line = content.split('\n')[1] if '\n' in content else content[:100]
            extracted_claim = f"PDF Text Excerpt: {first_line.strip()[:140]}"

        # 2. DOCX -> python-docx
        elif ext == 'docx':
            doc = docx.Document(file_path)
            lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables if present
            table_lines = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))
            
            all_parts = []
            if lines:
                all_parts.append("\n".join(lines))
            if table_lines:
                all_parts.append("\n[Tables]\n" + "\n".join(table_lines))
            
            content = "\n\n".join(all_parts) if all_parts else "(No text paragraphs or tables found in DOCX)"
            summary = f"python-docx extracted {len(lines)} paragraph(s) and {len(table_lines)} table row(s)."
            first_line = lines[0] if lines else "Document body parsed."
            extracted_claim = f"DOCX Excerpt: {first_line[:140]}"

        # 3. TXT -> normal text reading
        elif ext == 'txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            lines = [line.strip() for line in content.splitlines() if line.strip()]
            summary = f"Extracted {len(lines)} text lines ({len(content)} characters)."
            first_line = lines[0] if lines else "Plaintext content parsed."
            extracted_claim = f"Text Excerpt: {first_line[:140]}"

        # 4. CSV -> pandas
        elif ext == 'csv':
            df = pd.read_csv(file_path)
            content = df.to_string(index=False)
            num_rows, num_cols = df.shape
            summary = f"pandas extracted CSV table: {num_rows} rows × {num_cols} columns."
            cols_str = ", ".join(list(df.columns)[:5])
            extracted_claim = f"Tabular Schema: [{cols_str}] ({num_rows} rows)"

        # 5. XLSX -> pandas
        elif ext == 'xlsx':
            excel_file = pd.ExcelFile(file_path)
            sheet_sections = []
            total_rows = 0
            for sheet_name in excel_file.sheet_names:
                df = excel_file.parse(sheet_name)
                total_rows += len(df)
                sheet_sections.append(f"=== Sheet: {sheet_name} ({len(df)} rows) ===\n{df.to_string(index=False)}")
            content = "\n\n".join(sheet_sections)
            summary = f"pandas extracted {len(excel_file.sheet_names)} sheet(s) with {total_rows} total rows."
            extracted_claim = f"Excel Workbook: {len(excel_file.sheet_names)} sheet(s) indexed."

    except Exception as e:
        content = f"Error during extraction: {str(e)}"
        summary = f"Extraction failed: {str(e)}"
        extracted_claim = f"Extraction error on {filename}."

    return {
        'name': filename,
        'type': ext,
        'size': size_str,
        'summary': summary,
        'extracted_claim': extracted_claim,
        'content': content
    }
